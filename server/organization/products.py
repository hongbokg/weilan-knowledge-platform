"""Read raw workbook cells and existing parser tables. Never infer prices with LLM."""
import csv,hashlib,io,json,re,unicodedata
from decimal import Decimal,InvalidOperation
from bs4 import BeautifulSoup
def norm(s):return re.sub(r'[\s\-－—_]+','',unicodedata.normalize('NFKC',str(s))).upper()
def value(v):return '' if v is None else str(v).strip()
SENSITIVE=re.compile(r'密码|口令|身份证|工资|社保|银行账号|账户密码|审计报告|财务报表|password|passwd|\bpwd\b|secret|api[_ -]?key|access[_ -]?token',re.I)
ALIASES={'model':r'型号|model','brand':r'品牌|制造商|生产厂商','item':r'品名|产品名称|设备名称|货物名称|材料名称|名称','spec':r'规格|参数|技术要求|配置|说明','price':r'单价|单位价格','unit':r'^单位$|计量单位','category':r'类别|分类'}
def columns(rows):
 best=None
 for i,row in enumerate(rows[:35]):
  mapping={}
  for j,v in enumerate(row):
   s=value(v).replace('\n','')
   for field,pattern in ALIASES.items():
    if re.search(pattern,s,re.I) and field not in mapping:mapping[field]=j
  score=3*('model' in mapping)+3*('item' in mapping)+2*('price' in mapping)+('spec' in mapping)
  if score>=5 and (best is None or score>best[0]):best=(score,i,mapping)
 return best
def category(s):
 for name,pattern in [('摄像头','摄像|摄像机|IPC'),('交换机','交换机|switch'),('硬盘','硬盘|磁盘|HDD|SSD'),('道闸大门','道闸|闸机|大门'),('录像机','录像机|NVR|DVR'),('光网络','光纤|光模块|OLT|ONU'),('服务器','服务器|server'),('电源UPS','UPS|电源|蓄电池'),('空调','空调')]:
  if re.search(pattern,s,re.I):return name
 return '其他产品'
def parameters(spec):
 result={}
 for name,pattern,unit in [('像素',r'(\d+(?:\.\d+)?)\s*万(?:像素|高清|全彩|网络)?','万像素'),('端口数',r'(\d+)\s*(?:个)?(?:口|端口)','口'),('存储容量',r'(\d+(?:\.\d+)?)\s*(TB|GB)\b',''),('分辨率',r'(\d{3,5}\s*[×xX*]\s*\d{3,5})','')]:
  hits=list(re.finditer(pattern,spec,re.I))
  if len(hits)==1:
   m=hits[0];result[name]={'value':m.group(1),'unit':m.group(2).upper() if name=='存储容量' else unit,'evidence':m.group(0)}
 return result
def table_records(rows,sheet,path,kb,md5,modified,company='',project='',locations=None):
 detected=columns(rows)
 if not detected:return []
 _,start,mapping=detected;header=[value(x) for x in rows[start]];out=[]
 if any(SENSITIVE.search(value(v)) for row in rows for v in row):raise ValueError('restricted_content_requires_review')
 for offset,row in enumerate(rows[start+1:],start+2):
  def get(field):
   j=mapping.get(field);return value(row[j]) if j is not None and j<len(row) else ''
  item,model,brand=get('item'),get('model'),get('brand')
  if mapping.get('brand')==mapping.get('model') and model:
   m=re.match(r'^(海康威视|大华|华为|新华三|H3C|锐捷|希捷|西数|WD|TP-LINK)[\s/：:]*',model,re.I)
   if m:brand=m.group(1);model=model[m.end():].strip()
   else:brand=''
  if not (item or model) or re.search(r'^(合计|总计|小计|说明|备注|序号)$',item):continue
  if item in ('名称','设备名称','产品名称'):continue
  spec=get('spec');raw_price=get('price');price=''
  if raw_price:
   raw_price=re.sub(r'[,，￥¥\s]','',raw_price)
   if re.fullmatch(r'\d+(?:\.\d+)?',raw_price):
    try:
     num=Decimal(raw_price)
     if num>0:price=str(num)
    except InvalidOperation:pass
  price_header=header[mapping['price']] if 'price' in mapping else ''
  tax='未税' if re.search('不含税|未税',price_header) else '含税' if '含税' in price_header else '未知'
  unit=get('unit');currency='CNY' if re.search(r'人民币|[元￥¥]',price_header) else '未知'
  identity=hashlib.sha256(json.dumps([path,md5,sheet,offset],ensure_ascii=False).encode()).hexdigest()
  raw={'headers':header,'cells':[value(v) for v in row]}
  if locations:raw['location']=locations
  out.append(dict(id=identity,path=path,kb=kb,md5=md5,sheet=sheet,row=offset,item=item,brand=brand,model=model,model_key=norm(model),category=get('category') or category(' '.join([item,model,spec])),spec=spec,parameters=json.dumps(parameters(spec),ensure_ascii=False),price=price,tax=tax,unit=unit,currency=currency,modified=modified,company=company,project=project,raw=json.dumps(raw,ensure_ascii=False)))
 return out
def workbook(data,ext):
 if ext=='.xlsx':
  import openpyxl
  import zipfile
  with zipfile.ZipFile(io.BytesIO(data)) as archive:
   if sum(f.file_size for f in archive.infolist())>512*1024**2:raise ValueError('expanded_workbook_too_large')
  book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
  try:
   for sheet in book:
    if sheet.max_row and (sheet.max_row>100000 or sheet.max_column>200 or sheet.max_row*sheet.max_column>500000):raise ValueError('worksheet_too_large')
    rows=list(sheet.iter_rows(values_only=True))
    yield sheet.title,rows
  finally:book.close()
 elif ext=='.xls':
  import xlrd
  book=xlrd.open_workbook(file_contents=data,on_demand=True)
  try:
   for sheet in book.sheets():
    if sheet.nrows>100000:raise ValueError('worksheet_too_large')
    yield sheet.name,[sheet.row_values(i) for i in range(sheet.nrows)]
  finally:book.release_resources()
 else:
  try:text=data.decode('utf-8-sig')
  except UnicodeDecodeError:text=data.decode('gb18030')
  yield 'CSV',list(csv.reader(io.StringIO(text)))
def markdown_tables(text):
 soup=BeautifulSoup(text,'html.parser')
 for i,table in enumerate(soup.find_all('table'),1):
  grid=[];spans={}
  for rownum,tr in enumerate(table.find_all('tr')):
   row=[];col=0
   for cell in tr.find_all(['td','th'],recursive=False):
    while (rownum,col) in spans:row.append(spans[(rownum,col)]);col+=1
    t=cell.get_text(' ',strip=True);rs=min(int(cell.get('rowspan',1)),100);cs=min(int(cell.get('colspan',1)),100)
    for dc in range(cs):
     row.append(t)
     for dr in range(1,rs):spans[(rownum+dr,col+dc)]=t
    col+=cs
   while (rownum,col) in spans:row.append(spans[(rownum,col)]);col+=1
   grid.append(row)
  yield 'HTML表'+str(i),grid
 block=[];index=0
 for line in text.splitlines()+['']:
  if line.strip().startswith('|') and line.count('|')>=3:
   cells=[c.strip() for c in line.strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[:\-\s]+',c) for c in cells):block.append(cells)
  elif block:
   index+=1;yield 'Markdown表'+str(index),block;block=[]
